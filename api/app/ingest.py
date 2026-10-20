"""Carga el corpus preparado por el equipo en PostgreSQL: `python -m app.ingest [--sin-embeddings]`.

Lee de RAG_DIR: fuentes.csv (inventario), chunks.jsonl (fragmentos por artículo/sección) y datos/*.xlsx (lugares y aulas).
Es idempotente: vuelve a cargar todo lo que viene del corpus sin tocar los documentos subidos desde el portal.
"""

import csv
import json
import re
import sys
import time
import unicodedata

import openpyxl
import psycopg

from .catalog import TOPICS, acceso_for, audiencia_for, kind_for, topic_for, version_for
from .config import API_DIR, settings

RAG = settings.rag_dir


def slug(s: str) -> str:
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-")


def dms(s: str | None) -> tuple[float, float] | None:
    """0°25'03.9"N 78°11'17.5"W → (lat, lon)."""
    if not s:
        return None
    parts = re.findall(r"(\d+)°(\d+)'([\d.]+)\"?\s*([NSEW])", str(s))
    if len(parts) != 2:
        return None
    out = []
    for d, m, sec, h in parts:
        v = int(d) + int(m) / 60 + float(sec) / 3600
        out.append(-v if h in "SW" else v)
    return out[0], out[1]


# Nombres en inglés y alias para encontrar los lugares por su nombre común.
PLACE_EXTRA = {
    "Biblioteca Patrimonial": ("Heritage Library", ["biblioteca", "library", "sala de estudio", "study room"]),
    "Centro de Salud Tipo A": ("Health Center", ["centro de salud", "salud", "medico", "enfermeria", "health center", "doctor"]),
    "Krispan": ("Krispan minimarket", ["minimarket", "tienda", "krispan"]),
    "Containers": ("Containers food court", ["containers", "comida", "comedor", "food", "cafeteria"]),
    "Rectorado": ("Rector's office", ["rectorado", "rector"]),
    "Museo Arqueológico de YachayTech": ("Archaeological Museum", ["museo", "museum"]),
    "Jardín Botánico de YachayTech": ("Botanical Garden", ["jardin botanico", "botanical garden"]),
    "Hall Senescyt": ("Senescyt Hall", ["hall senescyt", "senescyt", "salon de eventos"]),
}


def load_places(c: psycopg.Connection) -> None:
    c.execute("DELETE FROM rooms; DELETE FROM places;")
    # Puntos importantes: nombre, uso, horario, coordenadas decimales y enlace de Google Maps.
    wb = openpyxl.load_workbook(RAG / "datos" / "puntos_importantes_universidad.xlsx", data_only=True)
    ws = wb.worksheets[0]
    n = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        name, uso, horario, coords, link = (list(row) + [None] * 5)[:5]
        if not name or not coords:
            continue
        lat, lon = (float(x) for x in str(coords).split(","))
        en, aliases = PLACE_EXTRA.get(str(name).strip(), (None, []))
        c.execute(
            "INSERT INTO places (id, name, name_en, uso, horario, lat, lon, maps_url, aliases, fuente) VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s) ON CONFLICT (id) DO NOTHING",
            (slug(str(name)), str(name).strip(), en, uso, horario, lat, lon, link, aliases, "puntos_importantes_universidad"),
        )
        n += 1
    # Distribución de aulas: zona / ubicación / aula / coordenadas en celdas combinadas (se rellenan hacia abajo).
    ws = openpyxl.load_workbook(RAG / "datos" / "Distribución Aulas_.xlsx", data_only=True).worksheets[0]
    zona = ubic = None
    coord: tuple[float, float] | None = None
    rooms = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        z, u, aula, co = (list(row) + [None] * 4)[:4]
        if z:
            zona = str(z).strip()
        if u:
            ubic = str(u).strip()
        if co:
            coord = dms(co)
        if not aula or not zona:
            continue
        code = str(aula).strip().split(" (")[0].strip()
        rooms.append((code, zona, ubic or zona, coord))
    zones: dict[str, tuple[float, float]] = {}
    for _, z, _, co in rooms:
        if co and z not in zones:
            zones[z] = co
    for z, (lat, lon) in zones.items():
        c.execute(
            "INSERT INTO places (id, name, name_en, uso, lat, lon, aliases, fuente) VALUES (%s,%s,%s,%s,%s,%s,%s,%s) ON CONFLICT (id) DO NOTHING",
            (f"zona-{slug(z)}", z, z, "Zona de aulas", lat, lon, [slug(z).replace("-", " ")], "distribucion_aulas"),
        )
    counts: dict[tuple[float, float], int] = {}
    for *_, co in rooms:
        if co:
            counts[co] = counts.get(co, 0) + 1
    for code, z, u, co in rooms:
        c.execute(
            "INSERT INTO rooms (code, zona, ubicacion, place_id, lat, lon, coord_compartida) VALUES (%s,%s,%s,%s,%s,%s,%s) ON CONFLICT (code) DO NOTHING",
            (code.upper(), z, u, f"zona-{slug(z)}", co[0] if co else None, co[1] if co else None, bool(co and counts[co] > 1)),
        )
    print(f"Lugares: {n} puntos + {len(zones)} zonas · aulas: {len(rooms)}")


def main() -> None:
    with_embeddings = "--sin-embeddings" not in sys.argv
    if not (RAG / "chunks.jsonl").exists():
        sys.exit(f"No encuentro el corpus en {RAG}. Revisa RAG_DIR en api/.env.")
    t0 = time.time()
    fuentes = [r for r in csv.DictReader(open(RAG / "fuentes.csv", encoding="utf-8")) if r.get("incluir") == "si"]
    chunks = [json.loads(line) for line in open(RAG / "chunks.jsonl", encoding="utf-8") if line.strip()]
    with psycopg.connect(settings.database_url) as c:
        for t in TOPICS:
            c.execute(
                "INSERT INTO topics (id, name, description, name_en, description_en) VALUES (%s,%s,%s,%s,%s) "
                "ON CONFLICT (id) DO UPDATE SET name=EXCLUDED.name, description=EXCLUDED.description, name_en=EXCLUDED.name_en, description_en=EXCLUDED.description_en",
                t,
            )
        c.execute("DELETE FROM chunks WHERE doc_id IN (SELECT id FROM documents WHERE uploaded_by = 'corpus')")
        for f in fuentes:
            doc_id = f["doc_id"]
            size = (RAG / "corpus" / f"{doc_id}.md").stat().st_size // 1024 if (RAG / "corpus" / f"{doc_id}.md").exists() else 0
            c.execute(
                """INSERT INTO documents (id, title, tipo, topic_id, audiencia, version, periodo, fecha_consulta, url, acceso, kind, size_kb, uploaded_by, notes)
                   VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,'corpus',%s)
                   ON CONFLICT (id) DO UPDATE SET title=EXCLUDED.title, tipo=EXCLUDED.tipo, topic_id=EXCLUDED.topic_id, audiencia=EXCLUDED.audiencia,
                     version=EXCLUDED.version, periodo=EXCLUDED.periodo, fecha_consulta=EXCLUDED.fecha_consulta, url=EXCLUDED.url,
                     acceso=EXCLUDED.acceso, kind=EXCLUDED.kind, size_kb=EXCLUDED.size_kb, notes=EXCLUDED.notes, status='indexado', is_current=true""",
                (
                    doc_id, f["titulo"], f["tipo"], topic_for(doc_id), audiencia_for(doc_id),
                    version_for(doc_id, f["titulo"], f.get("periodo_vigencia", "")), f.get("periodo_vigencia") or None,
                    f.get("fecha_consulta") or None, f.get("url") or None, acceso_for(f.get("fuente", ""), f.get("url", "")),
                    kind_for(f.get("archivo_origen", ""), f.get("url", "")), size, f.get("motivo") or None,
                ),
            )
        known = {f["doc_id"] for f in fuentes}
        rows = [ch for ch in chunks if ch["doc_id"] in known]
        vectors: list[str | None] = [None] * len(rows)
        if with_embeddings:
            from fastembed import TextEmbedding

            print(f"Calculando embeddings de {len(rows)} fragmentos con {settings.embedding_model} (la primera vez descarga el modelo)…")
            model = TextEmbedding(settings.embedding_model, cache_dir=str(API_DIR / ".cache"))
            texts = [f"{ch.get('seccion', '')}\n{ch['texto'][:2000]}" for ch in rows]
            vectors = ["[" + ",".join(f"{float(x):.6f}" for x in v) + "]" for v in model.embed(texts, batch_size=64)]
        with c.cursor() as cur:
            cur.executemany(
                "INSERT INTO chunks (id, doc_id, ord, seccion, texto, tokens, embedding) VALUES (%s,%s,%s,%s,%s,%s,%s::vector)",
                [
                    (ch["chunk_id"], ch["doc_id"], int(str(ch["chunk_id"]).rsplit("#", 1)[-1]) if "#" in str(ch["chunk_id"]) else i,
                     ch.get("seccion") or "", ch["texto"], int(ch.get("tokens_aprox") or 0), vectors[i])
                    for i, ch in enumerate(rows)
                ],
            )
        load_places(c)
        from .control_seed import seed_control

        seed_control(c)
        c.commit()
    print(f"Corpus cargado: {len(fuentes)} documentos, {len(rows)} fragmentos{' con embeddings' if with_embeddings else ''} en {time.time() - t0:.0f} s.")


if __name__ == "__main__":
    main()
