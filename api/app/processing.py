"""Procesa un documento subido desde el Portal: lo convierte a texto, lo parte en fragmentos, calcula embeddings y lo indexa.

Mismo criterio que el corpus inicial del equipo (docs/rag/rag/preparar_corpus.py): tablas del PDF intercaladas por posición,
sin membretes repetidos, encabezados por TÍTULO/CAPÍTULO/Art., fragmentos de ≤ 800 tokens con la ruta de sección como contexto.
Corre en segundo plano: el documento queda 'procesando' y pasa a 'indexado' (o 'error' con el motivo en notes).
"""

import csv
import io
import re
import traceback
import unicodedata
from collections import Counter
from pathlib import Path

from .config import API_DIR
from .db import conn

UPLOADS = API_DIR / "uploads"
MAX_TOK, MIN_TOK = 800, 120
tok = lambda s: len(s) // 4  # noqa: E731
nfc = lambda s: unicodedata.normalize("NFC", s)  # noqa: E731
ES_TABLA = lambda l: l.lstrip().startswith("|")  # noqa: E731
ES_ITEM = re.compile(r"^\s*([-*•·]|\d+[.)]|[a-zA-Z][.)]|\(\w\)|Art\.|Artículo|CAPÍTULO|TÍTULO|SECCIÓN|DISPOSICI|#)", re.I)


class ProcessingError(Exception):
    """Error explicable al administrador (se guarda en notes)."""


# ------------------------------------------------------------------ conversión

def _cells(l: str) -> list[str]:
    return [c.strip() for c in l.strip().strip("|").split("|")]


def _compact_table(lines: list[str]) -> list[str]:
    rows = [_cells(l) for l in lines if not re.fullmatch(r"\|?[\s:|-]+\|?", l.strip())]
    if not rows:
        return []
    n = max(map(len, rows))
    rows = [r + [""] * (n - len(r)) for r in rows]
    merged: list[list[str]] = []
    for r in rows:
        if merged and not r[0] and any(r):  # fila de continuación (celda combinada)
            for i, c in enumerate(r):
                if c:
                    merged[-1][i] = (merged[-1][i] + " " + c).strip()
        else:
            merged.append(r)
    empty = {i for i in range(n) if not any(r[i] for r in merged)}
    merged = [[c for i, c in enumerate(r) if i not in empty] for r in merged if any(r)]
    i = 0
    while merged and i < len(merged[0]) - 1:  # columnas partidas que nunca coinciden en una fila
        if not any(r[i] and r[i + 1] for r in merged):
            for r in merged:
                r[i] = r[i] or r[i + 1]
                del r[i + 1]
        else:
            i += 1
    if not merged or not merged[0]:
        return []
    return ["| " + " | ".join(merged[0]) + " |", "| " + " | ".join("---" for _ in merged[0]) + " |"] + ["| " + " | ".join(r) + " |" for r in merged[1:]]


def _render_table(rows) -> list[str]:
    rows = [[" ".join((c or "").split()) for c in r] for r in rows]
    if max(map(len, rows), default=0) <= 1:
        return [r[0] for r in rows if r and r[0]]
    return _compact_table(["| " + " | ".join(r) + " |" for r in rows])


def pdf_to_text(data: bytes) -> tuple[str, int]:
    import pdfplumber

    with pdfplumber.open(io.BytesIO(data)) as doc:
        n = len(doc.pages)
        pages, sigs = [], Counter()
        for p in doc.pages:
            items, y = [], 0.0
            for tb in sorted(p.find_tables(), key=lambda t: t.bbox[1]):
                top, bot = max(tb.bbox[1], 0), min(tb.bbox[3], p.height)
                if top - y > 2:
                    items.append(("t", p.crop((0, y, p.width, top)).extract_text() or ""))
                lines = _render_table(tb.extract())
                items.append(("tb", lines))
                sigs[re.sub(r"\d+", "#", " ".join(lines).lower())] += 1
                y = max(y, bot)
            if p.height - y > 2:
                items.append(("t", p.crop((0, y, p.width, p.height)).extract_text() or ""))
            pages.append(items)
    limit = max(3, int(n * 0.3))
    out = []
    for items in pages:
        for kind, c in items:
            if kind == "t":
                out.append(c)
            elif c and sigs[re.sub(r"\d+", "#", " ".join(c).lower())] < limit:  # tablas repetidas = membretes
                out.append("\n".join(c) if not c[0].startswith("|") else "\n" + "\n".join(c) + "\n")
    return "\n".join(out), n


def xlsx_to_text(data: bytes) -> tuple[str, int]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=False)
    parts = []
    for ws in wb.worksheets:
        merged = {}
        for rng in ws.merged_cells.ranges:  # celdas combinadas: el valor se repite en todas
            v = ws.cell(rng.min_row, rng.min_col).value
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    merged[(r, c)] = v
        rows = []
        for r in ws.iter_rows():
            vals = [merged.get((c.row, c.column), c.value) for c in r]
            vals = ["" if v is None else " ".join(str(v).split()) for v in vals]
            if any(vals):
                rows.append(vals)
        if rows:
            parts.append(f"## Hoja: {ws.title}\n\n" + "\n".join(_render_table(rows)))
    return "\n\n".join(parts), len(wb.worksheets)


def csv_to_text(data: bytes) -> tuple[str, int]:
    for enc in ("utf-8-sig", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = [[" ".join(c.split()) for c in r] for r in csv.reader(io.StringIO(text), dialect) if any(c.strip() for c in r)]
    return "\n".join(_render_table(rows)), 1


# ------------------------------------------------------------------ limpieza y estructura

def clean(text: str, pages: int = 0) -> str:
    text = nfc(text).replace("­", "").replace("\xa0", " ")
    lines = [l for l in text.splitlines() if not re.search(r"\.{5,}\s*\d+\s*$", l)]  # índice
    key = lambda l: re.sub(r"\d+", "#", re.sub(r"\s+", " ", l.strip().lower()))  # noqa: E731
    count = Counter(key(l) for l in lines if l.strip() and not ES_TABLA(l) and len(l.strip()) < 130)
    limit = max(5, int(pages * 0.3)) if pages else 10**9
    noise = {k for k, v in count.items() if v >= limit and not re.match(r"(art|#)", k)}  # encabezados y pies de página
    lines = [l for l in lines if key(l) not in noise and not re.fullmatch(r"\s*(pág\.?|página)?\s*\d+\s*(de|/)?\s*\d*\s*", l, re.I)]
    out, i = [], 0
    while i < len(lines):
        l = lines[i]
        if ES_TABLA(l):
            j = i
            while j < len(lines) and (ES_TABLA(lines[j]) or re.fullmatch(r"\s*\|?[\s:|-]+\|?\s*", lines[j])):
                j += 1
            out += _compact_table(lines[i:j]) + [""]
            i = j
            continue
        if not l.strip():
            out.append("")
        elif (out and out[-1] and not ES_TABLA(out[-1]) and not out[-1].startswith("#") and not ES_ITEM.match(l)
              and not re.search(r"[.:;]$", out[-1])) or (out and out[-1] and re.search(r"[a-záéíóú]-$", out[-1]) and l[:1].islower()):
            prev = out.pop()
            out.append((prev[:-1] if prev.endswith("-") and l[:1].islower() else prev + " ") + l.strip())
        else:
            out.append(l.rstrip())
        i += 1
    t = re.sub(r"[ \t]{2,}", " ", "\n".join(out))
    return re.sub(r"\n{3,}", "\n\n", t).strip()


def tipo_for(title: str) -> str:
    n = unicodedata.normalize("NFKD", title).encode("ascii", "ignore").decode().lower()
    for k, t in [("calendario", "calendario"), ("instructivo", "instructivo"), ("procedimiento", "procedimiento"), ("estatuto", "estatuto"),
                 ("reglamento", "reglamento"), ("ley", "ley"), ("codigo", "codigo"), ("plan", "plan"), ("programa", "programa"),
                 ("guia", "guia"), ("convocatoria", "convocatoria"), ("bases", "convocatoria")]:
        if k in n:
            return t
    return "documento"


def structure(text: str, tipo: str) -> str:
    res = []
    for l in text.split("\n"):
        s = l.strip()
        if tipo in ("ley", "reglamento", "codigo", "estatuto", "instructivo", "documento"):
            if re.match(r"^(TÍTULO|LIBRO)\s+[IVXLC\d]+", s) and len(s) < 120:
                l = "## " + s
            elif re.match(r"^(CAPÍTULO|SECCIÓN)\s+[IVXLC\d]+", s, re.I) and len(s) < 120:
                l = "### " + s
            elif re.match(r"^(Art\.|Artículo)\s*\d+", s):
                l = "#### " + s
        if tipo in ("procedimiento", "plan", "programa", "guia", "documento") and not s.startswith("|") and not l.startswith("#"):
            if re.match(r"^\d{1,2}(\.\d)?\.?\s+[A-ZÁÉÍÓÚÑ][A-ZÁÉÍÓÚÑ ,/Y]{3,70}$", s):
                l = "## " + s
        res.append(l)
    return "\n".join(res)


# ------------------------------------------------------------------ fragmentos

def _units(text: str):
    path = {2: "", 3: "", 4: ""}
    cur, buf = "", []
    for l in text.split("\n"):
        m = re.match(r"^(#{2,4}) (.*)", l)
        if m:
            if buf and "\n".join(buf).strip():
                yield cur, "\n".join(buf).strip()
            level = len(m.group(1))
            path[level] = m.group(2).strip()[:90]
            for k in path:
                if k > level:
                    path[k] = ""
            cur = " > ".join(v for k, v in sorted(path.items()) if v)
            buf = [l]
        else:
            buf.append(l)
    if buf and "\n".join(buf).strip():
        yield cur, "\n".join(buf).strip()


def _split(text: str) -> list[str]:
    if tok(text) <= MAX_TOK:
        return [text]
    blocks = [b for b in text.split("\n\n") if b.strip()]
    parts, cur, n = [], [], 0
    for b in blocks:
        if ES_TABLA(b.split("\n")[0]) and tok(b) > MAX_TOK:  # tabla grande: por filas, repitiendo el encabezado
            rows = b.split("\n")
            head, body = rows[:2], rows[2:]
            if cur:
                parts.append("\n\n".join(cur)); cur, n = [], 0
            g, gn = [], 0
            for r in body:
                if gn + tok(r) > MAX_TOK and g:
                    parts.append("\n".join(head + g)); g, gn = [], 0
                g.append(r); gn += tok(r)
            if g:
                parts.append("\n".join(head + g))
            continue
        if tok(b) > MAX_TOK:
            pieces = b.split("\n") if "\n" in b else re.split(r"(?<=[.;:])\s+", b)
            sub, sn = [], 0
            for p in pieces:
                if sn + tok(p) > MAX_TOK and sub:
                    parts.append(("\n" if "\n" in b else " ").join(sub)); sub, sn = [], 0
                sub.append(p); sn += tok(p)
            if sub:
                parts.append(("\n" if "\n" in b else " ").join(sub))
            continue
        if n + tok(b) > MAX_TOK and cur:
            parts.append("\n\n".join(cur))
            cur = [cur[-1]] if tok(cur[-1]) < 150 else []  # solape de un párrafo corto
            n = sum(map(tok, cur))
        cur.append(b); n += tok(b)
    if cur:
        parts.append("\n\n".join(cur))
    return parts


def make_chunks(doc_id: str, title: str, body: str) -> list[dict]:
    merged, acc = [], None
    for path, t in _units(body):
        acc = (acc[0], acc[1] + "\n\n" + t) if acc else (path, t)
        if tok(acc[1]) >= MIN_TOK:
            merged.append(acc); acc = None
    if acc:
        if merged:
            merged[-1] = (merged[-1][0], merged[-1][1] + "\n\n" + acc[1])
        else:
            merged.append(acc)
    out = []
    for path, t in merged:
        for p in _split(t):
            ctx = f"Documento: {title}" + (f" | Sección: {path}" if path else "")
            out.append({"id": f"{doc_id}#{len(out) + 1}", "seccion": path, "texto": f"{ctx}\n\n{p.strip()}", "tokens": tok(p) + tok(ctx)})
    return out


# ------------------------------------------------------------------ proceso completo

def extract(data: bytes, ext: str, title: str) -> tuple[str, int]:
    if ext == "pdf":
        raw, pages = pdf_to_text(data)
        if len(re.sub(r"\s", "", raw)) < 200:
            raise ProcessingError("El PDF no tiene texto seleccionable (parece una imagen escaneada). Súbelo con OCR o como PDF de texto.")
        return structure(clean(raw, pages), tipo_for(title)), pages
    if ext == "xlsx":
        raw, sheets = xlsx_to_text(data)
    elif ext == "csv":
        raw, sheets = csv_to_text(data)
    else:
        raise ProcessingError("Formato no admitido (solo PDF, XLSX o CSV).")
    if not raw.strip():
        raise ProcessingError("El archivo no tiene datos legibles.")
    return raw, sheets


def process(doc_id: str) -> None:
    """Procesa un documento en estado 'procesando'. Nunca lanza: deja 'indexado' o 'error' con el motivo."""
    from .search import vectors

    with conn() as c:
        d = c.execute("SELECT id, title, kind, replaces FROM documents WHERE id = %s", (doc_id,)).fetchone()
    if not d:
        return
    try:
        path = UPLOADS / f"{doc_id}.{d['kind']}"
        if not path.exists():
            raise ProcessingError("No se encontró el archivo subido en el servidor. Vuelve a subirlo.")
        text, pages = extract(path.read_bytes(), d["kind"], d["title"])
        chunks = make_chunks(doc_id, d["title"], text)
        if not chunks:
            raise ProcessingError("No se pudieron formar fragmentos con el contenido del archivo.")
        vecs = [vectors.embed(f"{ch['seccion']}\n{ch['texto'][:2000]}") for ch in chunks]
        with conn() as c:
            c.execute("DELETE FROM chunks WHERE doc_id = %s", (doc_id,))
            with c.cursor() as cur:
                cur.executemany(
                    "INSERT INTO chunks (id, doc_id, ord, seccion, texto, tokens, embedding) VALUES (%s,%s,%s,%s,%s,%s,%s::vector)",
                    [(ch["id"], doc_id, i + 1, ch["seccion"], ch["texto"], ch["tokens"], "[" + ",".join(f"{x:.6f}" for x in v) + "]")
                     for i, (ch, v) in enumerate(zip(chunks, vecs))],
                )
            unit = "páginas" if d["kind"] == "pdf" else "hojas" if d["kind"] == "xlsx" else "tabla"
            c.execute("UPDATE documents SET status = 'indexado', tipo = %s, notes = %s WHERE id = %s",
                      (tipo_for(d["title"]), f"{len(chunks)} fragmentos de {pages} {unit}." if unit != "tabla" else f"{len(chunks)} fragmentos.", doc_id))
            if d["replaces"]:  # la versión anterior deja de usarse solo cuando la nueva ya está indexada
                c.execute("UPDATE documents SET is_current = false WHERE id = %s", (d["replaces"],))
    except Exception as e:
        msg = str(e) if isinstance(e, ProcessingError) else f"Error al procesar el archivo: {type(e).__name__}. Revisa que no esté dañado o protegido."
        if not isinstance(e, ProcessingError):
            traceback.print_exc()
        with conn() as c:
            c.execute("UPDATE documents SET status = 'error', notes = %s WHERE id = %s", (msg, doc_id))


def pending() -> list[str]:
    """Documentos que quedaron 'procesando' (por ejemplo, si el servidor se reinició a mitad)."""
    with conn() as c:
        return [r["id"] for r in c.execute("SELECT id FROM documents WHERE status = 'procesando' AND uploaded_by <> 'corpus'").fetchall()]
